from flask import Flask, render_template, request, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///employee_referral.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db = SQLAlchemy(app)

class Team(db.Model):
    __tablename__ = "teams"
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)
    category = db.Column(db.String(50), nullable=False)
    actuarial_exam_required = db.Column(db.Boolean, default=False)
    employees = db.relationship("Employee", backref="team", lazy=True)

class Employee(db.Model):
    __tablename__ = "employees"
    id = db.Column(db.Integer, primary_key=True)
    employee_code = db.Column(db.String(50), unique=True, nullable=False)
    name = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(150), unique=True, nullable=False)
    team_id = db.Column(db.Integer, db.ForeignKey("teams.id"), nullable=False)
    joining_date = db.Column(db.Date)
    referrals = db.relationship("EmployeeReferral", backref="employee", lazy=True)

class JobRole(db.Model):
    __tablename__ = "job_roles"
    id = db.Column(db.Integer, primary_key=True)
    role_name = db.Column(db.String(150), nullable=False)
    category = db.Column(db.String(50), nullable=False)
    location = db.Column(db.String(100))
    actuarial_exam_required = db.Column(db.Boolean, default=False)
    actuarial_exam_optional = db.Column(db.Boolean, default=False)
    referrals = db.relationship("EmployeeReferral", backref="job_role", lazy=True)
    requirements = db.relationship("JobRequirement", backref="job_role", lazy=True,
                                    cascade="all, delete-orphan")

class JobRequirement(db.Model):
    __tablename__ = "job_requirements"
    id = db.Column(db.Integer, primary_key=True)
    job_role_id = db.Column(db.Integer, db.ForeignKey("job_roles.id"), nullable=False)
    requirement_type = db.Column(db.String(50), nullable=False)
    requirement_name = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text)
    mandatory = db.Column(db.Boolean, default=True)

class EmployeeReferral(db.Model):
    __tablename__ = "employee_referrals"
    id = db.Column(db.Integer, primary_key=True)
    employee_id = db.Column(db.Integer, db.ForeignKey("employees.id"), nullable=False)
    job_role_id = db.Column(db.Integer, db.ForeignKey("job_roles.id"), nullable=False)
    candidate_name = db.Column(db.String(150), nullable=False)
    candidate_email = db.Column(db.String(150), nullable=False)
    candidate_phone = db.Column(db.String(30))
    candidate_experience = db.Column(db.Integer)
    actuarial_exam = db.Column(db.String(150))
    referral_date = db.Column(db.Date, default=date.today)
    status = db.Column(db.String(50), default="Referred")

def seed_data():
    if Team.query.count() == 0:
        teams = [
            ("Technical Team 1", "Technical", False),
            ("Technical Team 2", "Technical", False),
            ("Technical Team 3", "Technical", False),
            ("Actuarial Team 1", "Actuarial", True),
            ("Actuarial Team 2", "Actuarial", True),
            ("Actuarial Team 3", "Actuarial", True),
            ("Actuarial Optional Team", "Actuarial", False),
            ("HR / Support", "HR / Support", False),
            ("IT", "IT", False),
        ]
        for name, category, required in teams:
            db.session.add(Team(name=name, category=category,
                                 actuarial_exam_required=required))
        db.session.commit()

    if JobRole.query.count() == 0:
        jobs = [
            ("Python Developer", "Technical", "Gurugram", False, False),
            ("Data Engineer", "Technical", "Gurugram", False, False),
            ("Actuarial Analyst", "Actuarial", "Gurugram", True, False),
            ("Actuarial Associate", "Actuarial", "Mumbai", True, False),
            ("Risk Analyst", "Actuarial", "Gurugram", False, True),
            ("HR Executive", "HR / Support", "Gurugram", False, False),
            ("IT Support Engineer", "IT", "Gurugram", False, False),
        ]
        for args in jobs:
            db.session.add(JobRole(
                role_name=args[0], category=args[1], location=args[2],
                actuarial_exam_required=args[3],
                actuarial_exam_optional=args[4]
            ))
        db.session.commit()

    if JobRequirement.query.count() == 0:
        jobs = {j.role_name: j for j in JobRole.query.all()}
        reqs = [
            ("Python Developer", "Skill", "Python", "Strong Python programming", True),
            ("Python Developer", "Skill", "Flask", "Experience with Flask", True),
            ("Python Developer", "Skill", "SQL", "Good SQL knowledge", True),
            ("Data Engineer", "Skill", "SQL", "Advanced SQL", True),
            ("Data Engineer", "Skill", "Python", "Python for data processing", True),
            ("Actuarial Analyst", "Exam", "Actuarial Exam", "Relevant actuarial examination", True),
            ("Actuarial Analyst", "Skill", "Excel", "Advanced Excel", True),
            ("Actuarial Associate", "Exam", "Actuarial Qualification", "Actuarial exam progress", True),
            ("Risk Analyst", "Exam", "Actuarial Exam", "Actuarial exam is beneficial but optional", False),
            ("HR Executive", "Skill", "Communication", "Strong communication skills", True),
            ("IT Support Engineer", "Skill", "IT Support", "IT troubleshooting experience", True),
        ]
        for job, typ, name, desc, mandatory in reqs:
            db.session.add(JobRequirement(
                job_role_id=jobs[job].id, requirement_type=typ,
                requirement_name=name, description=desc, mandatory=mandatory
            ))
        db.session.commit()

    if Employee.query.count() == 0:
        t1 = Team.query.filter_by(name="Technical Team 1").first()
        a1 = Team.query.filter_by(name="Actuarial Team 1").first()
        db.session.add_all([
            Employee(employee_code="EMP001", name="Rahul Sharma",
                     email="rahul@example.com", team_id=t1.id,
                     joining_date=date(2022, 5, 10)),
            Employee(employee_code="EMP002", name="Priya Singh",
                     email="priya@example.com", team_id=a1.id,
                     joining_date=date(2021, 8, 15)),
        ])
        db.session.commit()

with app.app_context():
    db.create_all()
    seed_data()

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/employees")
def get_employees():
    return jsonify([{
        "id": e.id, "employee_code": e.employee_code, "name": e.name,
        "email": e.email, "team": e.team.name,
        "team_category": e.team.category,
        "actuarial_exam_required": e.team.actuarial_exam_required
    } for e in Employee.query.all()])

@app.route("/api/jobs")
def get_jobs():
    result = []
    for j in JobRole.query.all():
        result.append({
            "id": j.id, "role_name": j.role_name, "category": j.category,
            "location": j.location,
            "actuarial_exam_required": j.actuarial_exam_required,
            "actuarial_exam_optional": j.actuarial_exam_optional,
            "requirements": [{
                "id": r.id, "type": r.requirement_type,
                "name": r.requirement_name, "description": r.description,
                "mandatory": r.mandatory
            } for r in j.requirements]
        })
    return jsonify(result)

@app.route("/api/referrals", methods=["POST"])
def create_referral():
    data = request.get_json() or {}
    employee = db.session.get(Employee, data.get("employee_id"))
    job = db.session.get(JobRole, data.get("job_role_id"))

    if not employee or not job:
        return jsonify(success=False, message="Invalid employee or job."), 400
    if not data.get("candidate_name"):
        return jsonify(success=False, message="Candidate name is required."), 400
    if not data.get("candidate_email"):
        return jsonify(success=False, message="Candidate email is required."), 400
    if job.actuarial_exam_required and not data.get("actuarial_exam"):
        return jsonify(success=False,
                       message="Actuarial exam information is required for this role."), 400

    referral = EmployeeReferral(
        employee_id=employee.id, job_role_id=job.id,
        candidate_name=data["candidate_name"],
        candidate_email=data["candidate_email"],
        candidate_phone=data.get("candidate_phone"),
        candidate_experience=data.get("candidate_experience") or None,
        actuarial_exam=data.get("actuarial_exam") or None,
        referral_date=date.today(), status="Referred"
    )
    db.session.add(referral)
    db.session.commit()
    return jsonify(success=True, message="Employee referral submitted successfully.",
                   referral_id=referral.id)

@app.route("/api/referrals")
def get_referrals():
    referrals = EmployeeReferral.query.order_by(EmployeeReferral.id.desc()).all()
    return jsonify([{
        "id": r.id, "employee": r.employee.name,
        "employee_code": r.employee.employee_code,
        "team": r.employee.team.name, "candidate_name": r.candidate_name,
        "candidate_email": r.candidate_email, "candidate_phone": r.candidate_phone,
        "job": r.job_role.role_name, "category": r.job_role.category,
        "experience": r.candidate_experience,
        "actuarial_exam": r.actuarial_exam,
        "referral_date": r.referral_date.strftime("%Y-%m-%d"),
        "status": r.status
    } for r in referrals])

@app.route("/export/referrals")
def export_referrals():
    referrals = EmployeeReferral.query.order_by(EmployeeReferral.id.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Referrals"
    headers = ["Referral ID", "Employee Code", "Employee Name", "Employee Team",
               "Candidate Name", "Candidate Email", "Candidate Phone", "Job Role",
               "Category", "Experience", "Actuarial Exam", "Referral Date", "Status"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    for r in referrals:
        ws.append([
            r.id, r.employee.employee_code, r.employee.name, r.employee.team.name,
            r.candidate_name, r.candidate_email, r.candidate_phone,
            r.job_role.role_name, r.job_role.category, r.candidate_experience,
            r.actuarial_exam, r.referral_date, r.status
        ])
        ws.cell(ws.max_row, 12).number_format = "yyyy-mm-dd"
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max((len(str(c.value)) if c.value else 0) for c in col) + 3, 40)
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True,
                     download_name="employee_referrals.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)
